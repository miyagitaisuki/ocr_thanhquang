"""Đường ống đầy đủ process_file với engine giả: PDF thật (tạo bằng Pillow), không cần model.
Kiểm: sinh .docx có chữ, cache bỏ qua OCR khi chạy lại, và cờ --no-cell-reocr được truyền xuống layout."""
import threading
from pathlib import Path

import pytest
from PIL import Image

from ocrnhat import engines, layout, pipeline
from ocrnhat.engines.base import OcrEngine
from ocrnhat.model import Block, Line, Page


class _CountingEngine(OcrEngine):
    calls = 0

    def __init__(self, device="cpu"):
        pass

    def ocr(self, img, name="page.png"):
        type(self).calls += 1
        h, w = img.shape[:2]
        return Page(w, h, [Block("text", 0, 0, w, h, [Line("東京です", x=10, y=10, w=200, h=30, conf=0.9)])])


@pytest.fixture
def fake_engine(monkeypatch):
    _CountingEngine.calls = 0
    monkeypatch.setattr(engines, "create", lambda name, device="cpu": _CountingEngine(device))
    monkeypatch.setattr(engines, "resolve", lambda name=None: (name if name not in (None, "auto") else "ndlocr", "test"))
    monkeypatch.setattr(engines, "engine_version", lambda name: "test")
    monkeypatch.setattr(pipeline, "_tls", threading.local())
    return _CountingEngine


@pytest.fixture
def pdf(tmp_path):
    path = tmp_path / "sample.pdf"
    Image.new("RGB", (800, 1000), "white").save(path, "PDF", resolution=200)
    return path


def test_process_file_sinh_docx_co_chu(fake_engine, pdf, tmp_path):
    out = tmp_path / "out"
    r = pipeline.process_file(pdf, out, 200, None, device="cpu", use_cache=False)
    assert r["pages"] == 1
    docx = Path(r["docx"])
    assert docx.exists()
    from docx import Document
    text = "\n".join(p.text for p in Document(docx).paragraphs)
    assert "東京です" in text


def test_cache_chay_lai_khong_goi_engine(fake_engine, pdf, tmp_path):
    out = tmp_path / "out"
    pipeline.process_file(pdf, out, 200, None, device="cpu", use_cache=True)
    first = fake_engine.calls
    assert first == 1
    pipeline.process_file(pdf, out, 200, None, device="cpu", use_cache=True)
    assert fake_engine.calls == first        # trang đã có trong cache: không OCR lại


def test_tat_cache_thi_ocr_lai(fake_engine, pdf, tmp_path):
    out = tmp_path / "out"
    pipeline.process_file(pdf, out, 200, None, device="cpu", use_cache=False)
    pipeline.process_file(pdf, out, 200, None, device="cpu", use_cache=False)
    assert fake_engine.calls == 2


@pytest.mark.parametrize("cell_reocr, expect_callable", [(True, True), (False, False)])
def test_cell_reocr_duoc_truyen_xuong_layout(fake_engine, pdf, tmp_path, monkeypatch, cell_reocr, expect_callable):
    seen = []
    real = layout.build_page

    def spy(page, img, reocr=None, kana_reader=None, lang="ja"):
        seen.append(reocr)
        return real(page, img, reocr, kana_reader, lang=lang)

    monkeypatch.setattr(layout, "build_page", spy)
    pipeline.process_file(pdf, tmp_path / "out", 200, None, device="cpu", use_cache=False, cell_reocr=cell_reocr)
    assert len(seen) == 1
    assert (seen[0] is not None) == expect_callable


def test_ket_qua_co_toa_do_dong_de_ra_soat(fake_engine, pdf, tmp_path):
    """Bản Qt rà soát trong cửa sổ xem trước: kết quả phải kèm tọa độ dòng, vùng cắt, phóng to và ngưỡng."""
    r = pipeline.process_file(pdf, tmp_path / "out", 200, None, device="cpu", use_cache=False,
                              crop=(10, 20, 500, 600), upscale=2.0, conf_threshold=0.7)
    assert r["crop"] == (10, 20, 500, 600) and r["upscale"] == 2.0 and r["conf_threshold"] == 0.7
    assert r["dpi"] == 200
    (page_no, lines), = r["page_lines"]
    assert page_no == 1
    x, y, x1, y1, conf, text = lines[0]
    assert text == "東京です" and conf == pytest.approx(0.9)
    assert x1 > x and y1 > y


def test_nhieu_vung_noi_ket_qua_va_doi_toa_do_ve_trang_goc(fake_engine, pdf, tmp_path):
    """Hai vùng: mỗi vùng OCR riêng; kết quả một trang nối theo thứ tự vùng; tọa độ dòng cộng vị trí vùng."""
    regions = [(10, 20, 500, 600), (300, 400, 700, 800)]
    r = pipeline.process_file(pdf, tmp_path / "out", 200, None, device="cpu", use_cache=False, crop=regions)
    assert r["pages"] == 1                       # đếm trang, không đếm vùng
    (page_no, lines), = r["page_lines"]
    assert page_no == 1 and len(lines) == 2
    xs = sorted((x, y) for x, y, *_ in lines)
    assert xs == [(10 + 10, 20 + 10), (300 + 10, 400 + 10)]   # fake engine trả dòng tại (10, 10) trong mỗi vùng
    from docx import Document
    text = "\n".join(p.text for p in Document(r["docx"]).paragraphs)
    assert text.count("東京です") == 2


def test_parse_crops_nhieu_vung():
    assert pipeline.parse_crops("1,2,30,40;50,60,70,80") == [(1, 2, 30, 40), (50, 60, 70, 80)]
    assert pipeline.parse_crops("") is None
    with pytest.raises(ValueError):
        pipeline.parse_crops("1,2,30,40;300,20,10,400")
