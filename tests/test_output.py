"""Test đầu ra: docx, txt, cache và đọc XML của NDLOCR (không cần model)."""
from docx import Document

from ocrnhat.cache import PageCache, file_hash
from ocrnhat.docx_out import write_docx, write_txt
from ocrnhat.engines.ndlocr import parse_page_xml
from ocrnhat.layout import Cell, Para, Table
from ocrnhat.model import Block, Line, Page


def _sample():
    t = Table(2, 2, [Cell(0, 0, rs=2, text="見出し"), Cell(0, 1, text="値1"),
                     Cell(1, 1, text="値2")])
    paras = [
        Para("本文です。", x=0, y=0, x1=10, y1=10, size=10),
        Para("見出し", kind="heading", x=0, y=0, x1=10, y1=10, size=14),
        t,
    ]
    return [(3, paras)]


def test_docx_co_doan_bang_va_tieu_de(tmp_path):
    out = tmp_path / "a.docx"
    write_docx(_sample(), out, page_marker=True)
    d = Document(out)
    texts = [p.text for p in d.paragraphs]
    assert any("【原本 p.3】" in t for t in texts)
    assert "本文です。" in texts
    assert "見出し" in texts
    assert len(d.tables) == 1
    tb = d.tables[0]
    assert len(tb.rows) == 2 and len(tb.columns) == 2
    assert tb.cell(0, 0).text == "見出し"   # ô gộp đọc được từ ô gốc
    assert tb.cell(1, 1).text == "値2"


def test_docx_ruby_la_phan_tu_cap_doan_va_khong_mat_chu(tmp_path):
    p = Para("日本語", x=0, y=0, x1=10, y1=10, size=10, ruby=[(0, 2, "にほんご")])
    out = tmp_path / "r.docx"
    write_docx([(1, [p])], out, ruby=True)
    body = Document(out).element.body
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    para = next(e for e in body.iter(W + "p") if e.find(f".//{W}ruby") is not None)
    ruby = para.find(W + "ruby")
    assert ruby is not None and ruby.getparent() is para        # không nằm trong w:r
    text = "".join(t.text for t in body.iter(W + "t"))
    assert "日本語" in text and "にほんご" in text


def test_txt_co_hang_bang(tmp_path):
    out = tmp_path / "a.txt"
    write_txt(_sample(), out)
    s = out.read_text(encoding="utf-8")
    assert "=== Trang 3 ===" in s
    assert "本文です。" in s
    assert "\t値1" in s


def test_cache_doc_lai_dung_trang(tmp_path):
    page = Page(100, 50, [Block("text", 1, 2, 3, 4, [Line("日本語", 1, 2, 3, 4, 0.9, "", 1)], "x")])
    c = PageCache(tmp_path / ".cache", "h", 200, "ndlocr", "v1")
    assert c.get(1) is None
    c.put(1, page)
    assert c.get(1) == page
    # khác engine hoặc phiên bản thì không dùng lại cache cũ
    assert PageCache(tmp_path / ".cache", "h", 200, "chandra", "v1").get(1) is None


def test_file_hash_thay_doi_theo_noi_dung(tmp_path):
    f = tmp_path / "x.bin"
    f.write_bytes(b"abc")
    h1 = file_hash(f)
    f.write_bytes(b"abd")
    assert file_hash(f) != h1


def test_parse_page_xml_doc_text_va_bang():
    xml = """<PAGE>
      <TEXTBLOCK><LINE STRING="日本語" X="10" Y="20" WIDTH="30" HEIGHT="40" CONF="0.9" ORDER="1"/></TEXTBLOCK>
      <BLOCK TYPE="表組" X="100" Y="200" WIDTH="300" HEIGHT="100">
        <LINE STRING="会社" X="110" Y="210" WIDTH="40" HEIGHT="40" CONF="0.8" ORDER="2"/>
      </BLOCK>
    </PAGE>"""
    page = parse_page_xml(xml, 800, 600)
    kinds = [b.kind for b in page.blocks]
    assert kinds == ["text", "table"]
    assert page.blocks[0].lines[0].text == "日本語"
    assert page.blocks[1].lines[0].conf == 0.8
    assert (page.width, page.height) == (800, 600)
